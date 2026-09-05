#!/usr/bin/env python3
"""Generate the constraint-kit training PowerPoint from markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "constraint-kit-training.md"
OUTPUT = ROOT / "UMN-SRE-constraint-kit-Training.pptx"

MAROON = RGBColor(87, 0, 20)
MAROON_DARK = RGBColor(66, 0, 15)
GOLD = RGBColor(255, 204, 51)
TEAL = RGBColor(0, 126, 121)
INK = RGBColor(28, 28, 30)
MUTED = RGBColor(92, 92, 98)
PAPER = RGBColor(250, 248, 244)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(224, 217, 210)

SERIF = "Georgia"
SANS = "Aptos"
MONO = "Menlo"


def parse_slides(markdown: str) -> list[dict[str, object]]:
    chunks = re.split(r"\n---\n", markdown)
    slides: list[dict[str, object]] = []
    for chunk in chunks:
        match = re.search(r"^## Slide (\d+):\s*(.+)$", chunk, re.MULTILINE)
        if not match:
            continue
        number = int(match.group(1))
        section_title = match.group(2).strip()
        title = section_title
        subtitle = ""
        bullets: list[str] = []
        notes = ""
        scenario = ""
        flow: list[str] = []
        exercise: list[str] = []

        title_match = re.search(r"^Title:\s*(.+)$", chunk, re.MULTILINE)
        if title_match:
            title = title_match.group(1).strip()
        subtitle_match = re.search(r"^Subtitle:\s*(.+?)(?:\n\n|\nBullets:)", chunk, re.MULTILINE | re.DOTALL)
        if subtitle_match:
            subtitle = " ".join(line.strip() for line in subtitle_match.group(1).splitlines()).strip()
        notes_match = re.search(r"Speaker notes:\n\n(.+)$", chunk, re.MULTILINE | re.DOTALL)
        if notes_match:
            notes = notes_match.group(1).strip()

        bullets_match = re.search(r"Bullets:\n\n(.+?)(?:\n\nSpeaker notes:|\n\nScenario:|\Z)", chunk, re.MULTILINE | re.DOTALL)
        if bullets_match:
            bullets = [line[2:].strip() for line in bullets_match.group(1).splitlines() if line.startswith("- ")]

        scenario_match = re.search(r"Scenario:\n\n(.+?)(?:\n\nFlow:|\n\nSpeaker notes:)", chunk, re.MULTILINE | re.DOTALL)
        if scenario_match:
            scenario = " ".join(line.strip() for line in scenario_match.group(1).splitlines()).strip()
        flow_match = re.search(r"Flow:\n\n(.+?)(?:\n\nSpeaker notes:|\Z)", chunk, re.MULTILINE | re.DOTALL)
        if flow_match:
            flow = [re.sub(r"^\d+\.\s*", "", line).strip() for line in flow_match.group(1).splitlines() if re.match(r"^\d+\.\s*", line)]
        exercise_match = re.search(r"Exercise:\n\n(.+?)(?:\n\nSpeaker notes:|\Z)", chunk, re.MULTILINE | re.DOTALL)
        if exercise_match:
            exercise = [re.sub(r"^\d+\.\s*", "", line).strip() for line in exercise_match.group(1).splitlines() if re.match(r"^\d+\.\s*", line)]

        slides.append(
            {
                "number": number,
                "section_title": section_title,
                "title": title,
                "subtitle": subtitle,
                "bullets": bullets,
                "notes": notes,
                "scenario": scenario,
                "flow": flow,
                "exercise": exercise,
            }
        )
    return slides


def fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def set_text(box, text: str, font: str, size: int, color: RGBColor, bold: bool = False, italic: bool = False) -> None:
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color
    paragraph.font.bold = bold
    paragraph.font.italic = italic


def add_text(slide, left, top, width, height, text, font=SANS, size=18, color=INK, bold=False, italic=False, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    set_text(box, text, font, size, color, bold, italic)
    if align is not None:
        box.text_frame.paragraphs[0].alignment = align
    return box


def add_footer(slide, number: int) -> None:
    add_text(slide, Inches(0.62), Inches(7.08), Inches(0.55), Inches(0.22), "umn-sre", SANS, 9, MUTED, bold=True)
    add_text(slide, Inches(1.2), Inches(7.08), Inches(2.2), Inches(0.22), "constraint-kit training", SANS, 9, MUTED)
    add_text(slide, Inches(12.0), Inches(7.08), Inches(0.7), Inches(0.22), f"{number} / 29", SANS, 9, MUTED, align=PP_ALIGN.RIGHT)


def add_eyebrow(slide, text: str) -> None:
    add_text(slide, Inches(0.72), Inches(0.42), Inches(3.8), Inches(0.28), text.upper(), SANS, 8, MAROON, bold=True)


def add_title(slide, title: str, y: float = 0.75, size: int = 27) -> None:
    add_text(slide, Inches(0.72), Inches(y), Inches(8.4), Inches(0.65), title, SERIF, size, INK, bold=True)


def add_notes(slide, notes: str) -> None:
    if not notes:
        return
    notes_tf = slide.notes_slide.notes_text_frame
    notes_tf.text = notes


def add_bullet_panel(slide, bullets: list[str], left: float, top: float, width: float, height: float, accent: RGBColor = MAROON) -> None:
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    fill(panel, WHITE)
    panel.line.color.rgb = LINE
    panel.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left), Inches(top), Inches(0.06), Inches(height))
    fill(bar, accent)
    box = slide.shapes.add_textbox(Inches(left + 0.28), Inches(top + 0.22), Inches(width - 0.45), Inches(height - 0.35))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0)
    frame.margin_right = Inches(0)
    frame.margin_top = Inches(0)
    frame.margin_bottom = Inches(0)
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"• {bullet}"
        paragraph.level = 0
        paragraph.font.name = SANS
        paragraph.font.size = Pt(14 if len(bullets) < 6 else 12)
        paragraph.font.color.rgb = INK
        paragraph.space_after = Pt(8 if len(bullets) < 6 else 5)


def add_step_cards(slide, items: list[str], top: float = 2.15) -> None:
    count = len(items)
    columns = 3 if count > 4 else count
    card_w = 3.75 if columns == 3 else 11.2 / columns
    card_h = 1.0 if count > 4 else 1.35
    gap = 0.35
    for index, item in enumerate(items):
        row = index // columns
        col = index % columns
        left = 0.72 + col * (card_w + gap)
        y = top + row * (card_h + 0.35)
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(y), Inches(card_w), Inches(card_h))
        fill(card, WHITE)
        card.line.color.rgb = LINE
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(left + 0.18), Inches(y + 0.2), Inches(0.38), Inches(0.38))
        fill(circle, MAROON)
        num = add_text(slide, Inches(left + 0.18), Inches(y + 0.25), Inches(0.38), Inches(0.22), str(index + 1), SANS, 9, WHITE, bold=True, align=PP_ALIGN.CENTER)
        num.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        add_text(slide, Inches(left + 0.72), Inches(y + 0.18), Inches(card_w - 0.9), Inches(card_h - 0.2), item, SANS, 12, INK, bold=False)


def add_title_slide(prs: Presentation, slide_data: dict[str, object]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(0.55), Inches(12.25), Inches(6.15))
    fill(bg, MAROON)
    dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(10.7), Inches(1.1), Inches(0.95), Inches(0.95))
    fill(dot, GOLD)
    ring = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.78), Inches(5.95), Inches(0.55), Inches(0.55))
    ring.fill.background()
    ring.line.color.rgb = TEAL
    ring.line.width = Pt(2.5)
    add_text(slide, Inches(1.18), Inches(1.65), Inches(4.2), Inches(0.35), "UMN-SRE TEAM TRAINING", SANS, 10, GOLD, bold=True)
    add_text(slide, Inches(1.18), Inches(2.25), Inches(9.5), Inches(0.9), str(slide_data["title"]), SERIF, 32, WHITE, bold=True)
    add_text(slide, Inches(1.18), Inches(3.08), Inches(9.2), Inches(0.52), str(slide_data["subtitle"]), SERIF, 18, PAPER, italic=True)
    bullets = slide_data["bullets"]
    if isinstance(bullets, list):
        add_text(slide, Inches(1.18), Inches(4.08), Inches(9.2), Inches(0.75), " · ".join(bullets[:3]), SANS, 11, PAPER)
    add_text(slide, Inches(1.18), Inches(5.65), Inches(4.8), Inches(0.28), "github.com/umn-sre  ·  VS Code + GitHub Copilot", MONO, 9, GOLD)
    add_notes(slide, str(slide_data["notes"]))


def add_standard_slide(prs: Presentation, slide_data: dict[str, object]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = PAPER
    number = int(slide_data["number"])
    add_eyebrow(slide, str(slide_data["section_title"]))
    add_title(slide, str(slide_data["title"]))
    subtitle = str(slide_data.get("subtitle") or "")
    if subtitle:
        add_text(slide, Inches(0.74), Inches(1.34), Inches(10.8), Inches(0.42), subtitle, SERIF, 15, MUTED, italic=True)
    bullets = slide_data.get("bullets")
    if isinstance(bullets, list) and bullets:
        if number in {5, 7, 22}:
            add_step_cards(slide, bullets, top=2.05)
        elif number in {6, 15, 17, 24, 25, 26, 27}:
            add_step_cards(slide, bullets, top=2.0)
        else:
            add_bullet_panel(slide, bullets, 0.72, 2.0, 11.35, 3.95)

    scenario = str(slide_data.get("scenario") or "")
    flow = slide_data.get("flow")
    exercise = slide_data.get("exercise")
    if scenario:
        scenario_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(1.85), Inches(11.35), Inches(0.72))
        fill(scenario_box, MAROON)
        add_text(slide, Inches(1.0), Inches(2.08), Inches(10.7), Inches(0.28), scenario, SANS, 13, WHITE, bold=True)
    if isinstance(flow, list) and flow:
        add_step_cards(slide, flow, top=2.9)
    if isinstance(exercise, list) and exercise:
        add_step_cards(slide, exercise, top=2.0)

    if number in {3, 8, 18, 19, 20, 28}:
        accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(11.1), Inches(0.65), Inches(0.6), Inches(0.6))
        fill(accent, GOLD)
    if number in {6, 15, 17}:
        line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.72), Inches(1.76), Inches(10.8), Inches(0.03))
        fill(line, MAROON)
    add_footer(slide, number)
    add_notes(slide, str(slide_data["notes"]))


def build() -> None:
    slides = parse_slides(SOURCE.read_text())
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for slide_data in slides:
        if slide_data["number"] == 1:
            add_title_slide(prs, slide_data)
        else:
            add_standard_slide(prs, slide_data)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
